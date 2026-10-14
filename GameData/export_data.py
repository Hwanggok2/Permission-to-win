"""Validate the MVP Excel source and export the three Unity JSON v1 files."""
from __future__ import annotations

import argparse
import json
import math
import os
from pathlib import Path
import re
import sys
import tempfile
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / "GameData_MVP.xlsx"
GENERATED = ROOT / "Generated"
INCOMPLETE_MARKER = ".export-incomplete"
SCHEMA_VERSION = 1
SCHEMAS = {
    "Stages": [
        ("ExportType", "string", "All"), ("Id", "int", "All"),
        ("Chapter", "int", "All"), ("TitleKey", "string", "Client"),
        ("AiDepth", "int", "All"), ("AiNodeBudget", "int", "All"),
        ("AiRandomPermille", "int", "All"), ("Starter", "int", "All"),
        ("OpeningMoves", "string", "All"),
        ("LearningGoal", "none", "None"), ("VerificationNote", "none", "None"),
    ],
    "AdConfig": [
        ("ExportType", "string", "All"),
        ("InterstitialGraceMatches", "int", "Client"),
        ("InterstitialEveryMatches", "int", "Client"),
        ("InterstitialCooldownSeconds", "int", "Client"),
        ("BannerLobby", "bool", "Client"), ("BannerResult", "bool", "Client"),
        ("RewardedHintsPerMatch", "int", "Client"),
    ],
    "Strings": [
        ("ExportType", "string", "All"), ("SearchKey", "string", "Client"),
        ("Content_Kor", "string", "Client"), ("/Description", "none", "None"),
        ("Content_Eng", "string", "Client"),
    ],
}


class ValidationError(ValueError):
    pass


def fail(location: str, message: str):
    raise ValidationError(f"{location}: {message}")


def typed_value(value, kind: str, location: str):
    if value is None:
        return None
    if kind in ("string", "none"):
        if not isinstance(value, str):
            fail(location, "expected a text cell")
        return value.strip()
    if kind == "bool":
        if type(value) is not bool:
            fail(location, "expected an Excel TRUE/FALSE cell")
        return value
    if kind == "int":
        if type(value) not in (int, float) or not math.isfinite(value):
            fail(location, "expected a finite numeric cell")
        if value != int(value) or not -(2**31) <= value <= 2**31 - 1:
            fail(location, "expected an Int32 integer")
        return int(value)
    fail(location, f"unsupported type {kind!r}")


def read_table(sheet, schema):
    if sheet.merged_cells.ranges:
        fail(sheet.title, "merged cells are not allowed in an import sheet")
    if sheet.tables:
        fail(sheet.title, "use a plain range with four metadata rows, not Excel Table objects")
    if sheet.max_column != len(schema):
        fail(sheet.title, f"expected exactly {len(schema)} columns")
    for col, (name, kind, scope) in enumerate(schema, 1):
        for row, expected in enumerate((name, kind, scope), 1):
            cell = sheet.cell(row, col)
            if cell.value != expected:
                fail(f"{sheet.title}!{cell.coordinate}", f"expected {expected!r}")
        description = sheet.cell(4, col)
        if not isinstance(description.value, str) or not description.value.strip():
            fail(f"{sheet.title}!{description.coordinate}", "Korean field description is required")
    rows = []
    for cells in sheet.iter_rows():
        for cell in cells:
            if cell.data_type in ("f", "e"):
                fail(f"{sheet.title}!{cell.coordinate}", "formula/error cells are forbidden")
        if cells[0].row <= 4 or all(c.value is None for c in cells):
            continue
        row_number = cells[0].row
        export = cells[0].value
        if export not in ("USE", "NONE"):
            fail(f"{sheet.title}!A{row_number}", "ExportType must be USE or NONE")
        if export == "NONE":
            continue
        row = {"_row": row_number}
        for cell, (name, kind, _) in zip(cells, schema):
            row[name] = typed_value(cell.value, kind, f"{sheet.title}!{cell.coordinate}")
        rows.append(row)
    return rows


def read_source(path: Path):
    workbook = load_workbook(path, data_only=False, read_only=False, keep_links=False)
    try:
        if set(workbook.sheetnames) != set(SCHEMAS):
            fail(path.name, "expected exactly Stages, AdConfig, Strings sheets")
        return {name: read_table(workbook[name], schema) for name, schema in SCHEMAS.items()}
    finally:
        workbook.close()


def integer(row, name, low, high, location):
    value = row.get(name)
    if type(value) is not int or not low <= value <= high:
        fail(location, f"{name} must be an integer from {low} to {high}")
    return value


def required_text(row, name, location):
    value = row.get(name)
    if not isinstance(value, str) or not value.strip():
        fail(location, f"{name} is required")
    return value.strip()


def format_indices(text: str, location: str):
    """Validate the supported C# format subset: {0}, {1}, and escaped braces."""
    indices, offset = set(), 0
    while offset < len(text):
        token = re.match(r"\{\{|\}\}|\{(0|[1-9][0-9]*)\}|[^{}]+", text[offset:])
        if token is None:
            fail(location, f"invalid format placeholder/braces at character {offset + 1}")
        if token.group(1) is not None:
            indices.add(int(token.group(1)))
        offset += token.end()
    if sorted(indices) != list(range(len(indices))):
        fail(location, "format indices must be consecutive from 0")
    return indices


def has_four(board, col, row, player):
    for dx, dy in ((1, 0), (0, 1), (1, 1), (1, -1)):
        count = 1
        for sign in (-1, 1):
            x, y = col + sign * dx, row + sign * dy
            while 0 <= x < 7 and 0 <= y < 6 and board[y][x] == player:
                count += 1
                x, y = x + sign * dx, y + sign * dy
        if count >= 4:
            return True
    return False


def replay_opening(moves: str, starter: int, location="openingMoves"):
    """Replay every prefix. Return bottom-up board and next side; never accept a terminal prefix."""
    if type(starter) is not int or starter not in (-1, 1):
        fail(location, "starter must be 1 or -1")
    if not isinstance(moves, str) or (moves and not re.fullmatch(r"[0-6](,[0-6])*", moves)):
        fail(location, "use comma-separated columns 0..6, without spaces, or empty text")
    columns = [int(value) for value in moves.split(",")] if moves else []
    board, heights = [[0] * 7 for _ in range(6)], [0] * 7
    side = starter
    for turn, col in enumerate(columns, 1):
        row = heights[col]
        if row == 6:
            fail(location, f"move {turn}: column {col} is full")
        board[row][col] = side
        heights[col] += 1
        if has_four(board, col, row, side):
            fail(location, f"move {turn}: opening already contains a victory")
        if turn == 42:
            fail(location, "move 42: opening is already a draw")
        side = -side
    return board, side


def can_survive_next_reply(board, side):
    """Existence of a win/draw or a move that prevents every immediate opponent win."""
    heights = [next((y for y in range(6) if board[y][x] == 0), 6) for x in range(7)]
    for col, row in enumerate(heights):
        if row == 6:
            continue
        after = [line[:] for line in board]
        after[row][col] = side
        if has_four(after, col, row, side):
            return True
        opponent_wins = False
        for reply in range(7):
            landing = heights[reply] + (reply == col)
            if landing == 6:
                continue
            after[landing][reply] = -side
            opponent_wins |= has_four(after, reply, landing, -side)
            after[landing][reply] = 0
        if not opponent_wins:
            return True
    return False


def validate(tables):
    strings, keys = [], set()
    for row in tables["Strings"]:
        location = f"Strings row {row['_row']}"
        key = required_text(row, "SearchKey", location)
        if not re.fullmatch(r"[A-Za-z][A-Za-z0-9_]*", key):
            fail(location, "SearchKey must start with an ASCII letter and use letters/digits/underscore")
        if key in keys:
            fail(location, f"duplicate SearchKey {key!r}")
        keys.add(key)
        ko = required_text(row, "Content_Kor", location)
        en = required_text(row, "Content_Eng", location)
        ko_indices = format_indices(ko, location + " Content_Kor")
        en_indices = format_indices(en, location + " Content_Eng")
        if ko_indices != en_indices:
            fail(location, f"format indices differ between ko {sorted(ko_indices)} and en {sorted(en_indices)}")
        strings.append({"key": key, "ko": ko, "en": en})
    if not strings:
        fail("Strings", "at least one active string is required")

    stages, ids = [], set()
    for row in tables["Stages"]:
        location = f"Stages row {row['_row']}"
        stage_id = integer(row, "Id", 1, 30, location)
        if stage_id in ids:
            fail(location, f"duplicate Id {stage_id}")
        ids.add(stage_id)
        chapter = integer(row, "Chapter", 1, 3, location)
        if chapter != (stage_id - 1) // 10 + 1:
            fail(location, "chapter must contain its consecutive ten stage IDs")
        key = required_text(row, "TitleKey", location)
        if key not in keys:
            fail(location, f"TitleKey {key!r} is missing from active Strings")
        starter = integer(row, "Starter", -1, 1, location)
        moves = row.get("OpeningMoves")
        moves = "" if moves is None else moves
        board, side = replay_opening(moves, starter, location + " OpeningMoves")
        if not can_survive_next_reply(board, side):
            fail(location, "every legal move loses on the opponent's next reply")
        stages.append({
            "id": stage_id, "chapter": chapter, "titleKey": key,
            "aiDepth": integer(row, "AiDepth", 1, 6, location),
            "aiNodeBudget": integer(row, "AiNodeBudget", 100, 30000, location),
            "aiRandomPermille": integer(row, "AiRandomPermille", 0, 1000, location),
            "starter": starter, "openingMoves": moves,
        })
        required_text(row, "LearningGoal", location)
        note = required_text(row, "VerificationNote", location)
        if note != "강제승리 미검증":
            fail(location, "VerificationNote must remain '강제승리 미검증' until separately proven")
    if ids != set(range(1, 31)):
        fail("Stages", "active IDs must be exactly 1..30 (three chapters of ten)")

    if len(tables["AdConfig"]) != 1:
        fail("AdConfig", "exactly one active configuration is required")
    row = tables["AdConfig"][0]
    location = f"AdConfig row {row['_row']}"
    # Contract v1 defines these product trial values, not a platform policy.
    expected = {
        "InterstitialGraceMatches": 3, "InterstitialEveryMatches": 2,
        "InterstitialCooldownSeconds": 90, "BannerLobby": True,
        "BannerResult": True, "RewardedHintsPerMatch": 1,
    }
    ads = {"schemaVersion": SCHEMA_VERSION}
    for field, value in expected.items():
        if type(row.get(field)) is not type(value) or row[field] != value:
            fail(location, f"{field} must be {value!r} under contract v1")
        ads[field[0].lower() + field[1:]] = value
    return {
        "stages.json": {"schemaVersion": SCHEMA_VERSION, "stages": sorted(stages, key=lambda s: s["id"])},
        "ads.json": ads,
        "strings.json": {"schemaVersion": SCHEMA_VERSION, "strings": sorted(strings, key=lambda s: s["key"])},
    }


def serialize(payloads):
    return {name: (json.dumps(value, ensure_ascii=False, indent=2, allow_nan=False) + "\n").encode("utf-8")
            for name, value in payloads.items()}


def write_outputs(payloads):
    """Stage the bundle; restore previous outputs on handled replacement failures.

    A retained marker blocks --check after interruption or failed rollback. This is
    not a filesystem transaction or a guarantee against power loss/process kills.
    """
    files = serialize(payloads)
    GENERATED.mkdir(exist_ok=True)
    marker = GENERATED / INCOMPLETE_MARKER
    if marker.exists():
        raise OSError(f"Incomplete export: inspect {marker} and its backup directory before retrying")
    pending = Path(tempfile.mkdtemp(prefix=".export-", dir=GENERATED))
    previous, replaced = {}, []
    marker_owned = False
    try:
        for name, content in files.items():
            (pending / name).write_bytes(content)
            destination = GENERATED / name
            previous[name] = destination.exists()
            if previous[name]:
                (pending / ("previous-" + name)).write_bytes(destination.read_bytes())
        # Exclusive creation also stops another exporter from publishing concurrently.
        with marker.open("x", encoding="utf-8") as stream:
            marker_owned = True
            json.dump({"backupDirectory": pending.name, "previouslyPresent": previous}, stream)
            stream.write("\n")
        try:
            for name in files:
                os.replace(pending / name, GENERATED / name)
                replaced.append(name)
            marker.unlink()
        except BaseException as failure:
            rollback_errors = []
            for name in reversed(replaced):
                try:
                    if previous[name]:
                        os.replace(pending / ("previous-" + name), GENERATED / name)
                    else:
                        (GENERATED / name).unlink(missing_ok=True)
                except OSError as error:
                    rollback_errors.append(f"{name}: {error}")
            if rollback_errors:
                raise OSError(f"Export failed and rollback is incomplete. Keep {marker} and {pending}; "
                              + "; ".join(rollback_errors)) from failure
            # Every changed destination has its prior contents/existence restored.
            marker.unlink(missing_ok=True)
            raise
    finally:
        if not (marker_owned and marker.exists()):
            # Only the six known files in our unique directory; no recursive removal.
            for name in files:
                (pending / name).unlink(missing_ok=True)
                (pending / ("previous-" + name)).unlink(missing_ok=True)
            pending.rmdir()


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source", type=Path, default=SOURCE)
    parser.add_argument("--check", action="store_true", help="validate and compare generated files without writing")
    args = parser.parse_args(argv)
    try:
        if (GENERATED / INCOMPLETE_MARKER).exists():
            fail("Generated", "incomplete export marker exists; restore/verify the saved bundle before retrying")
        payloads = validate(read_source(args.source))
        if args.check:
            for name, content in serialize(payloads).items():
                path = GENERATED / name
                if not path.exists() or path.read_bytes() != content:
                    fail(name, "missing or stale generated file; run exporter")
        else:
            write_outputs(payloads)
    except (ValidationError, OSError, ValueError, KeyError, BadZipFile, InvalidFileException) as error:
        print(f"GameData validation failed: {args.source}: {error}", file=sys.stderr)
        return 1
    print(f"Validated 30 stages / 3 chapters, {len(payloads['strings.json']['strings'])} strings, ads v1; "
          + ("generated files match." if args.check else "exported to GameData/Generated."))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
