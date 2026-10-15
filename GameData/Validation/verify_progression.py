"""Check the approved source edit, unchanged content and deterministic export."""
from copy import copy
import hashlib
import json
from pathlib import Path
import sys
from openpyxl import load_workbook

ROOT=Path(__file__).resolve().parents[2]
sys.path.insert(0,str(ROOT/'GameData'))
import export_data

def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest().upper()

before=load_workbook(ROOT/'Client/Archive/Progression-20261015/GameData_MVP-before.xlsx')
after=load_workbook(ROOT/'GameData/GameData_MVP.xlsx')
assert before.sheetnames==after.sheetnames
changes={}
for name in before.sheetnames:
    old,new=before[name],after[name]
    assert old.freeze_panes==new.freeze_panes
    assert [str(v) for v in old.data_validations.dataValidation]==[str(v) for v in new.data_validations.dataValidation]
    assert {k:v.height for k,v in old.row_dimensions.items()}=={k:new.row_dimensions[k].height for k in old.row_dimensions}
    assert {k:v.width for k,v in old.column_dimensions.items()}=={k:new.column_dimensions[k].width for k in old.column_dimensions}
    changed=[]
    for row in old:
        for previous in row:
            current=new[previous.coordinate]
            for field in ('font','fill','border','alignment','protection'):
                assert copy(getattr(previous,field))==copy(getattr(current,field)),(name,previous.coordinate,field)
            if name!='Stages' or previous.column!=14 or previous.row<5:
                assert previous.number_format==current.number_format
            if previous.value!=current.value:
                if name=='Stages':
                    assert previous.coordinate in ('L4','N2','N4') or (5<=previous.row<=34 and previous.column in (5,6,7,12,14,15,16,17))
                elif name=='Strings':
                    assert previous.column in (3,5)
                    key=old.cell(previous.row,2).value
                    assert key.startswith('stage_') or key.startswith('rank_') or key.startswith('ai_rank_10_'),key
                else:
                    raise AssertionError('Unrelated AdConfig edit')
                changed.append(previous.coordinate)
    changes[name]=changed
before.close();after.close()
payloads=export_data.validate(export_data.read_source(export_data.SOURCE))
oldstages=json.loads((ROOT/'Client/Archive/Progression-20261015/stages-before.json').read_text(encoding='utf-8'))['stages']
stages=payloads['stages.json']['stages']
for old,new in zip(oldstages,stages):
    for key in ('id','chapter','titleKey','openingMoves','starter','descriptionKey'):
        assert old[key]==new[key],(new['id'],key)
assert [s['difficultyStars'] for s in stages]==[1+i*.5 for i in range(15) for _ in range(2)]
assert len({s['rankKey'] for s in stages})==18
assert max(sum(s['rankKey']==key for s in stages) for key in {s['rankKey'] for s in stages})==2
strings={r['key']:r for r in payloads['strings.json']['strings']}
assert strings[stages[-1]['rankKey']]['ko']=='사단장'
assert strings[stages[-1]['rankKey']]['en']=='Division Commander'
assert len(strings)==269
for name in payloads:
    generated=ROOT/'GameData/Generated'/name
    assert generated.read_bytes()==(ROOT/'Client/Assets/StageGame/Resources/GameData'/name).read_bytes()
    assert json.loads(generated.read_text(encoding='utf-8'))==payloads[name]
report={'status':'PASS','sourceSha256':sha(export_data.SOURCE),'stages':30,'rankTitles':18,'maxStagesPerRank':2,'starValues':[1+i*.5 for i in range(15)],'strings':len(strings),'preserved':'All30IDs,chapters,titleKeys,openings,starter,descriptions; AdConfig; original cell styles except star decimal format; original dimensions,panes,validation.','changedOriginalCells':changes,'generatedSha256':{n:sha(ROOT/'GameData/Generated'/n) for n in payloads}}
path=ROOT/'GameData/Validation/progression-20261015/verification.json'
path.write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
print('PASS:30stages,18ranks,1..8half-stars; preserved content/styles; Generated and Client copies match.')
