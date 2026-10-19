import hashlib
import json
import sys
from copy import copy
from pathlib import Path
from zipfile import ZipFile
import openpyxl

work=Path(__file__).parent
root=work.parents[2]
before=openpyxl.load_workbook(work/'before.xlsx')
after=openpyxl.load_workbook(work/'after.xlsx')
assert before.sheetnames==after.sheetnames
differences=[]
for name in before.sheetnames:
    a,b=before[name],after[name]
    assert b.max_column==a.max_column,(name,'columns')
    assert b.max_row==a.max_row+(7 if name=='Strings' else 0),(name,'rows')
    for row in a:
        for c in row:
            d=b[c.coordinate]
            for attr in ('value','data_type','number_format','alignment','font','fill','border','protection','hyperlink','comment'):
                av,bv=copy(getattr(c,attr)),copy(getattr(d,attr))
                if av!=bv: differences.append((name,c.coordinate,attr,str(av),str(bv)))
    for attr in ('freeze_panes','merged_cells','print_area','print_title_cols','print_title_rows','sheet_format','sheet_properties','sheet_view','sheet_state','auto_filter','data_validations'):
        av,bv=getattr(a,attr),getattr(b,attr)
        if av!=bv: differences.append((name,attr,str(av),str(bv)))
    for dimtype in ('column_dimensions','row_dimensions'):
        av,bv=getattr(a,dimtype),getattr(b,dimtype)
        for key,value in av.items():
            if dict(value)!=dict(bv[key]): differences.append((name,dimtype,str(key),str(dict(value)),str(dict(bv[key]))))
with ZipFile(work/'before.xlsx') as a,ZipFile(work/'after.xlsx') as b:
    names=set(a.namelist())|set(b.namelist())
    changed=[n for n in sorted(names) if (a.read(n) if n in a.namelist() else None)!=(b.read(n) if n in b.namelist() else None)]
report={'originalRows':{s.title:s.max_row for s in before},'newRows':{s.title:s.max_row for s in after},'unrelatedDifferences':differences,'changedZipParts':changed}
(work/'preservation.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
assert not differences,'Unrelated workbook properties changed'
print(json.dumps({'unrelatedWorkbookDifferences':len(differences),'changedZipParts':changed},ensure_ascii=False))
if '--outputs' not in sys.argv: sys.exit(0)
prior={row['key']:row for row in json.loads((work/'strings-before.json').read_text(encoding='utf-8'))['strings']}
final={row['key']:row for row in json.loads((root/'GameData/Generated/strings.json').read_text(encoding='utf-8'))['strings']}
added=sorted(set(final)-set(prior))
removed=sorted(set(prior)-set(final))
changed=sorted(key for key in prior if prior[key]!=final.get(key))
expected=json.loads((work/'changes.json').read_text(encoding='utf-8'))['changes']
assert added==sorted(x[0] for x in expected)
assert not removed and not changed,(removed,changed)
for key,ko,en in expected:
    assert final[key]=={'key':key,'ko':ko,'en':en},key
hashes=json.loads((work/'unchanged-hashes-before.json').read_text(encoding='utf-8-sig'))
for entry in hashes:
    assert hashlib.sha256(Path(entry['Path']).read_bytes()).hexdigest().upper()==entry['Hash'],entry['Path']
assert (root/'GameData/GameData_MVP.xlsx').read_bytes()==(work/'after.xlsx').read_bytes()
assert (root/'GameData/Generated/strings.json').read_bytes()==(root/'Client/Assets/StageGame/Resources/GameData/strings.json').read_bytes()
report={'stringCount':len(final),'addedKeys':added,'changedExistingKeys':changed,'removedKeys':removed,'runtimeStringsIdentical':True,'otherGeneratedFilesUnchanged':[x['Path'] for x in hashes]}
(work/'output-verification.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(report,ensure_ascii=False))
