"""Contract, invalid-input, reproducibility, and independent board checks."""
import copy
import hashlib
import io
import json
from contextlib import redirect_stderr, redirect_stdout
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from openpyxl import load_workbook

import export_data as data
import audit_openings as audit


DRAW = "5,3,3,2,6,2,1,4,4,2,6,3,0,4,2,5,4,2,2,6,0,1,3,4,5,3,4,6,1,1,6,0,6,5,5,0,1,5,3,1,0,0"


def independent_board(sequence, starter):
    """Column stacks plus a full 4-cell window scan, separate from the production last-move scan."""
    stacks = [[] for _ in range(7)]
    side = starter
    for step, column in enumerate(sequence, 1):
        if len(stacks[column]) >= 6:
            raise AssertionError("full column")
        stacks[column].append(side)
        cells = {(x, y): stone for x, stack in enumerate(stacks) for y, stone in enumerate(stack)}
        for x in range(7):
            for y in range(6):
                for dx, dy in ((1, 0), (0, 1), (1, 1), (-1, 1)):
                    line = [cells.get((x + n * dx, y + n * dy)) for n in range(4)]
                    if line[0] in (1, -1) and line == [line[0]] * 4:
                        raise AssertionError(f"terminal prefix {step}")
        if step == 42:
            raise AssertionError("draw")
        side = -side
    return stacks, side


def immediate_wins(stacks, side):
    result = []
    for column, stack in enumerate(stacks):
        if len(stack) >= 6:
            continue
        cells = {(x, y): stone for x, s in enumerate(stacks) for y, stone in enumerate(s)}
        cells[column, len(stack)] = side
        if any(all(cells.get((x + n * dx, y + n * dy)) == side for n in range(4))
               for x in range(7) for y in range(6)
               for dx, dy in ((1, 0), (0, 1), (1, 1), (-1, 1))):
            result.append(column)
    return result


class ExportTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.original = data.read_source(data.SOURCE)

    def setUp(self):
        self.tables = copy.deepcopy(self.original)

    def assert_invalid(self, part, index, field, value, message):
        self.tables[part][index][field] = value
        with self.assertRaisesRegex(data.ValidationError, message):
            data.validate(self.tables)

    def test_real_source_matches_exact_contract(self):
        payloads = data.validate(self.tables)
        stages = payloads["stages.json"]["stages"]
        self.assertEqual(30, len(stages))
        self.assertEqual(509, len(payloads["strings.json"]["strings"]))
        for entry in payloads["strings.json"]["strings"]:
            self.assertEqual({"key", "ko", "en"}, set(entry))
        expected = {"id", "chapter", "titleKey", "aiDepth", "aiNodeBudget", "aiRandomPermille", "starter", "openingMoves",
                    "rankKey", "descriptionKey", "difficultyStars", "aiIntroKey", "aiWinKey", "aiLoseKey", "characterId", "aiRating", "aiOversightPermille"}
        for stage in stages:
            self.assertEqual(expected, set(stage))
        self.assertEqual([10, 10, 10], [sum(s["chapter"] == c for s in stages) for c in (1, 2, 3)])

    def test_rank_and_star_boundaries(self):
        stages = data.validate(self.tables)['stages.json']['stages']
        self.assertEqual([f'C{n:02}' for n in list(range(1,13))+[n for n in range(13,19) for _ in range(2)]+[2,3,4,5,19,1]], [s['characterId'] for s in stages])
        self.assertEqual([2] * 15, [sum(s['difficultyStars'] == 1+n*.5 for s in stages) for n in range(15)])
        for index, field, value, message in ((2, 'RankKey', 'rank_02', 'RankKey must'),
                                             (3, 'RankKey', 'rank_01', 'RankKey must'),
                                             (5, 'DifficultyStars', 1, 'once per two'),
                                             (6, 'DifficultyStars', 1, 'once per two')):
            with self.subTest(index=index, field=field):
                self.tables = copy.deepcopy(self.original)
                self.assert_invalid('Stages', index, field, value, message)

    def test_multiplayer_required_unique_keys_and_bounds(self):
        cases = [('missing', 'missing configuration key'), ('duplicate', 'duplicate'),
                 ('unknown', 'unknown'), ('initial', 'initialRating'), ('fraction', 'eloK'),
                 ('search', 'exceeds maximum'), ('rank', 'missing rank string')]
        for case, message in cases:
            with self.subTest(case=case):
                tables = copy.deepcopy(self.original)
                entries = {r['Key']: r for r in tables['Multiplayer']}
                if case == 'missing': tables['Multiplayer'].pop()
                elif case == 'duplicate': tables['Multiplayer'].append(copy.deepcopy(tables['Multiplayer'][0]))
                elif case == 'unknown': tables['Multiplayer'][0]['Key'] = 'unsupported'
                elif case == 'initial': entries['initialRating']['Value'] = 300
                elif case == 'fraction': entries['eloK']['Value'] = 31.5
                elif case == 'search':
                    entries['searchStart']['Value'] = 300
                    entries['searchMaximum']['Value'] = 100
                elif case == 'rank': tables['Strings'] = [r for r in tables['Strings'] if r['SearchKey'] != 'mp_rank_20']
                with self.assertRaisesRegex(data.ValidationError, message): data.validate(tables)

    def test_missing_presentation_references(self):
        for field in ('RankKey', 'DescriptionKey', 'AiIntroKey', 'AiWinKey', 'AiLoseKey'):
            with self.subTest(field=field):
                self.tables = copy.deepcopy(self.original)
                self.assert_invalid('Stages', 0, field, 'missing', 'missing from active Strings')

    def test_ai_dialogue_roles_and_stage_descriptions(self):
        for field, value in (('AiWinKey', 'ai_rank_01_lose'), ('AiLoseKey', 'ai_rank_01_win'),
                             ('AiIntroKey', 'ai_rank_02_intro'), ('DescriptionKey', 'stage_desc_002')):
            with self.subTest(field=field):
                self.tables = copy.deepcopy(self.original)
                self.assert_invalid('Stages', 0, field, value, f'{field} must')

    def test_stars_require_bounded_half_steps(self):
        for value in (0, 8.5, 1.25, float('nan'), float('inf'), True, '1', None):
            with self.subTest(value=value):
                self.assert_invalid('Stages', 0, 'DifficultyStars', value, 'half-step')

    def test_ai_settings_cannot_reverse_the_difficulty_curve(self):
        for field,value in (('AiDepth',3),('AiNodeBudget',5000),('AiRandomPermille',300)):
            with self.subTest(field=field):
                self.tables=copy.deepcopy(self.original)
                self.assert_invalid('Stages',29,field,value,'same AI profile|must not decrease')

    def test_rating_profiles_and_oversight_are_validated(self):
        stages=data.validate(self.tables)['stages.json']['stages']
        self.assertEqual([200+100*(n//2) for n in range(30)],[s['aiRating'] for s in stages])
        self.assertGreater(stages[0]['aiOversightPermille'],0)
        self.assertLessEqual(stages[0]['aiOversightPermille'],400)
        self.assertTrue(all(s['aiOversightPermille']==0 and s['aiDepth']>=2 for s in stages if s['aiRating']>=400))
        self.assertEqual(0,stages[-1]['aiOversightPermille'])
        for field,value in (('AiRating',199),('AiRating',1601),('AiRating',250),('AiOversightPermille',-1),('AiOversightPermille',1001)):
            with self.subTest(field=field,value=value):
                self.tables=copy.deepcopy(self.original)
                self.assert_invalid('Stages',0,field,value,'AiRating|AiOversightPermille')
        self.tables=copy.deepcopy(self.original)
        self.assert_invalid('Stages',1,'AiOversightPermille',0,'same AI profile')

    def test_presentation_copy_is_one_line_without_arguments(self):
        for key in ('story_opponent_001', 'story_stage_desc_001', 'story_002_intro', 'story_002_win', 'story_002_lose'):
            for value in ('First\nSecond', 'Line {0}'):
                with self.subTest(key=key, value=value):
                    self.tables = copy.deepcopy(self.original)
                    row = next(r for r in self.tables['Strings'] if r['SearchKey'] == key)
                    row.update(Content_Kor=value, Content_Eng=value)
                    with self.assertRaisesRegex(data.ValidationError, 'one line without format arguments'):
                        data.validate(self.tables)

    def test_balance_preserves_stage_identity_with_empty_boards_and_ai_starts(self):
        record = json.loads((data.ROOT / 'Validation/rank-ui-changes.json').read_text(encoding='utf-8'))
        stages = data.validate(self.tables)['stages.json']['stages']
        for previous, current in zip(record['previousStageFields'], stages):
            for key in ('id','chapter','titleKey'):
                self.assertEqual(previous[key],current[key])
        self.assertEqual(24, sum(not stage['openingMoves'] for stage in stages))
        self.assertEqual([5,8,11,14,18,22,26,30], [s['id'] for s in stages if s['starter'] == -1])
        self.assertGreaterEqual(sum(not s['openingMoves'] for s in stages[10:24]), 10)
        for stage in stages:
            self.assertLessEqual(len(stage['openingMoves'].split(',')) if stage['openingMoves'] else 0, 4)
        for index in (3,10):
            self.assertEqual('', stages[index]['openingMoves'], 'Stages 4/11 must no longer be one-turn puzzles')

    def test_first_stage_story_roles_and_exact_user_copy(self):
        payloads = data.validate(self.tables)
        stages = payloads['stages.json']['stages']
        strings = {row['key']: row for row in payloads['strings.json']['strings']}
        for field, suffix in (('aiIntroKey', 'intro'), ('aiWinKey', 'win'), ('aiLoseKey', 'lose')):
            self.assertEqual(f'story_first_{suffix}', stages[0][field])
            self.assertEqual(f'story_002_{suffix}', stages[1][field])
            self.assertEqual(f'story_003_{suffix}', stages[2][field])
        self.assertEqual('이 게임의 프로라고 했지? 그런데 내가 이길 것 같아.', strings['story_first_intro']['ko'])
        self.assertEqual('프로라는 말, 거짓말이었나?', strings['story_first_win']['ko'])
        self.assertEqual('프로는 역시 다르네. 다음엔 내가 이긴다.', strings['story_first_lose']['ko'])

    def test_first_story_exception_is_limited_to_stage_one_and_correct_role(self):
        for field, suffix in (('AiIntroKey', 'intro'), ('AiWinKey', 'win'), ('AiLoseKey', 'lose')):
            with self.subTest(field=field):
                self.tables = copy.deepcopy(self.original)
                self.assert_invalid('Stages', 1, field, f'story_first_{suffix}', f'{field} must')
        self.tables = copy.deepcopy(self.original)
        self.assert_invalid('Stages', 0, 'AiWinKey', 'story_first_lose', 'AiWinKey must')

    def test_all_openings_have_independent_legal_nonterminal_prefixes(self):
        for stage in self.tables["Stages"]:
            with self.subTest(stage=stage["Id"]):
                sequence = [int(x) for x in (stage["OpeningMoves"] or "").split(",") if x]
                stacks, side = independent_board(sequence, stage["Starter"])
                self.assertIn(side, (1,-1))
                actual, next_side = data.replay_opening(stage["OpeningMoves"] or "", stage["Starter"])
                self.assertEqual(side, next_side)
                self.assertEqual(sum(map(len, stacks)), sum(v != 0 for r in actual for v in r))

    def test_learning_motifs_and_difficulty_variety(self):
        themes = {row["LearningGoal"] for row in self.tables["Stages"]}
        self.assertGreaterEqual(len(themes), 4)
        depths = [row["AiDepth"] for row in self.tables["Stages"]]
        self.assertTrue(all(b >= a for a, b in zip(depths, depths[1:])))
        budgets=[row['AiNodeBudget'] for row in self.tables['Stages']]
        rates=[row['AiRandomPermille'] for row in self.tables['Stages']]
        self.assertTrue(all(b>=a for a,b in zip(budgets,budgets[1:])))
        self.assertTrue(all(b<=a for a,b in zip(rates,rates[1:])))
        for stage in self.tables["Stages"]:
            with self.subTest(stage=stage["Id"]):
                sequence = [int(x) for x in (stage["OpeningMoves"] or "").split(",") if x]
                stacks, side = independent_board(sequence, stage["Starter"])
                self.assertEqual([], immediate_wins(stacks, side))
                self.assertEqual([], immediate_wins(stacks, -side))

    def test_opening_grammar(self):
        for value in ("7", "-1", "1,", ",1", "1,,2", "1, 2", "0.0", "０", None):
            with self.subTest(value=value), self.assertRaises(data.ValidationError):
                data.replay_opening(value, 1)

    def test_starter_validation(self):
        for starter in (0, 2, True, "1"):
            with self.subTest(starter=starter), self.assertRaises(data.ValidationError):
                data.replay_opening("", starter)
        self.assertEqual(-1, data.replay_opening("", -1)[1])
        self.assertEqual(1, data.replay_opening("3", -1)[1])

    def test_full_column_and_six_legal_stones(self):
        data.replay_opening("0,0,0,0,0,0", 1)
        with self.assertRaisesRegex(data.ValidationError, "column 0 is full"):
            data.replay_opening("0,0,0,0,0,0,0", 1)

    def test_wins_and_moves_after_a_win_are_rejected(self):
        for sequence in ("0,1,0,1,0,1,0", "0,0,1,1,2,2,3", "0,1,0,1,0,1,0,2"):
            with self.subTest(sequence=sequence), self.assertRaisesRegex(data.ValidationError, "victory"):
                data.replay_opening(sequence, 1)

    def test_both_diagonal_directions_and_nonwin(self):
        for coordinates in ([(0, 0), (1, 1), (2, 2), (3, 3)], [(0, 3), (1, 2), (2, 1), (3, 0)]):
            board = [[0] * 7 for _ in range(6)]
            for x, y in coordinates:
                board[y][x] = -1
            x, y = coordinates[2]
            self.assertTrue(data.has_four(board, x, y, -1))
            board[coordinates[0][1]][coordinates[0][0]] = 1
            self.assertFalse(data.has_four(board, x, y, -1))

    def test_draw_prefix_and_draw_rejection(self):
        prefix = DRAW.rsplit(",", 1)[0]
        independent_board([int(x) for x in prefix.split(",")], 1)
        data.replay_opening(prefix, 1)
        with self.assertRaisesRegex(data.ValidationError, "draw"):
            data.replay_opening(DRAW, 1)

    def test_duplicate_stage(self):
        self.assert_invalid("Stages", 1, "Id", 1, "duplicate Id")

    def test_missing_stage(self):
        self.tables["Stages"].pop()
        with self.assertRaisesRegex(data.ValidationError, "exactly 1..30"):
            data.validate(self.tables)

    def test_chapter_membership(self):
        self.assert_invalid("Stages", 0, "Chapter", 2, "consecutive ten")

    def test_numeric_contract_boundaries(self):
        for field, values in {"AiDepth": [0, 7], "AiNodeBudget": [99, 30001], "AiRandomPermille": [-1, 1001], "Id": [0, 31], "Chapter": [0, 4]}.items():
            for value in values:
                with self.subTest(field=field, value=value):
                    self.tables = copy.deepcopy(self.original)
                    self.assert_invalid("Stages", 0, field, value, "integer")

    def test_missing_string_reference(self):
        self.assert_invalid("Stages", 0, "TitleKey", "missing", "missing from active Strings")

    def test_duplicate_and_blank_strings(self):
        self.assert_invalid("Strings", 1, "SearchKey", "stage_001", "duplicate SearchKey")
        self.tables = copy.deepcopy(self.original)
        self.assert_invalid("Strings", 0, "Content_Kor", " ", "required")

    def test_missing_english_is_rejected(self):
        for value in (None, "", " \n "):
            with self.subTest(value=value):
                self.assert_invalid("Strings", 0, "Content_Eng", value, "Content_Eng is required")

    def test_format_indices_allow_reordering_and_repetition(self):
        row = self.tables["Strings"][0]
        row.update(Content_Kor="{0} 다음 {1}, 다시 {0}", Content_Eng="{1} after {0}")
        data.validate(self.tables)

    def test_literal_braces_are_not_arguments(self):
        self.assertEqual(set(), data.format_indices("{{0}} {{name}}", "test"))
        self.assertEqual({0}, data.format_indices("{{{0}}}", "test"))

    def test_translation_argument_mismatch_is_rejected(self):
        row = self.tables["Strings"][0]
        for ko, en in (("{0}", "Text"), ("Text", "{0}"), ("{0} {1}", "{0}")):
            with self.subTest(ko=ko, en=en):
                row.update(Content_Kor=ko, Content_Eng=en)
                with self.assertRaisesRegex(data.ValidationError, "format indices differ"):
                    data.validate(self.tables)

    def test_malformed_or_unsupported_format_is_rejected_in_either_language(self):
        for field in ("Content_Kor", "Content_Eng"):
            for value in ("{", "}", "{0", "0}", "{{0}", "{name}", "{}", "{-1}",
                          "{00}", "{0:N2}", "{0,3}", "{0!r}", "{ 0}", "{０}"):
                with self.subTest(field=field, value=value):
                    self.tables = copy.deepcopy(self.original)
                    self.assert_invalid("Strings", 0, field, value, "invalid format")

    def test_argument_indices_must_start_at_zero_and_have_no_gaps(self):
        for value in ("{1}", "{0} {2}", "{999999999999999999}"):
            with self.subTest(value=value), self.assertRaisesRegex(data.ValidationError, "consecutive from 0"):
                data.format_indices(value, "test")

    def test_missing_english_column_metadata_is_rejected(self):
        workbook = load_workbook(data.SOURCE)
        try:
            sheet = workbook["Strings"]
            sheet["E1"] = "Content_English"
            with self.assertRaisesRegex(data.ValidationError, "Strings!E1"):
                data.read_table(sheet, data.SCHEMAS["Strings"])
        finally:
            workbook.close()

    def test_bad_translation_preserves_generated_bundle(self):
        with tempfile.TemporaryDirectory(prefix=".test-", dir=data.GENERATED) as directory:
            with patch.object(data, "GENERATED", Path(directory)), redirect_stdout(io.StringIO()), redirect_stderr(io.StringIO()):
                self.assertEqual(0, data.main([]))
                before = {p.name: p.read_bytes() for p in Path(directory).glob("*.json")}
                self.tables["Strings"][0]["Content_Eng"] = "Unclosed {0"
                with patch.object(data, "read_source", return_value=self.tables):
                    self.assertEqual(1, data.main([]))
                self.assertEqual(before, {p.name: p.read_bytes() for p in Path(directory).glob("*.json")})

    def test_unproven_note_required(self):
        self.assert_invalid("Stages", 0, "VerificationNote", "승리 보장", "separately proven")

    def test_ads_contract_and_row_count(self):
        self.assertEqual({
            "schemaVersion": 1, "interstitialGraceMatches": 0,
            "interstitialEveryMatches": 1, "interstitialCooldownSeconds": 0,
            "bannerLobby": True, "bannerResult": True, "rewardedHintsPerMatch": 1,
        }, data.validate(self.tables)["ads.json"])
        for field, value in [
            ("InterstitialGraceMatches", 3), ("InterstitialEveryMatches", 2),
            ("InterstitialCooldownSeconds", 90), ("InterstitialGraceMatches", 6),
            ("InterstitialCooldownSeconds", 180), ("InterstitialGraceMatches", -1),
            ("InterstitialEveryMatches", 0), ("InterstitialCooldownSeconds", True),
            ("BannerLobby", False), ("BannerResult", False), ("RewardedHintsPerMatch", 2),
        ]:
            with self.subTest(field=field, value=value):
                self.tables = copy.deepcopy(self.original)
                self.assert_invalid("AdConfig", 0, field, value, "contract v1")
        self.tables = copy.deepcopy(self.original)
        self.tables["AdConfig"].append(copy.deepcopy(self.tables["AdConfig"][0]))
        with self.assertRaisesRegex(data.ValidationError, "exactly one"):
            data.validate(self.tables)

    def test_cell_types_blank_zero_and_finite(self):
        self.assertEqual(0, data.typed_value(0, "int", "test"))
        self.assertIsNone(data.typed_value(None, "int", "test"))
        for value, kind in [(True, "int"), (1.5, "int"), (float("nan"), "int"), (float("inf"), "int"), (2**31, "int"), ("3", "int"), (1, "bool"), (1, "string")]:
            with self.subTest(value=value, kind=kind), self.assertRaises(data.ValidationError):
                data.typed_value(value, kind, "test")

    def test_metadata_and_export_filter_without_saving_workbook(self):
        workbook = load_workbook(data.SOURCE)
        try:
            sheet = workbook["Stages"]
            sheet["B2"] = "long"
            with self.assertRaisesRegex(data.ValidationError, "Stages!B2"):
                data.read_table(sheet, data.SCHEMAS["Stages"])
            sheet["B2"] = "int"
            sheet["A34"] = "NONE"
            self.assertEqual(29, len(data.read_table(sheet, data.SCHEMAS["Stages"])))
            sheet["I34"] = "=1+1"
            with self.assertRaisesRegex(data.ValidationError, "formula/error"):
                data.read_table(sheet, data.SCHEMAS["Stages"])
            sheet["I34"] = ""
            sheet["A34"] = ""
            with self.assertRaisesRegex(data.ValidationError, "ExportType"):
                data.read_table(sheet, data.SCHEMAS["Stages"])
        finally:
            workbook.close()

    def test_output_bytes_are_independent_of_row_order(self):
        expected = data.serialize(data.validate(self.tables))
        self.tables["Stages"].reverse()
        self.tables["Strings"].reverse()
        self.assertEqual(expected, data.serialize(data.validate(self.tables)))

    def test_repeated_export_matches_and_bad_input_preserves_output(self):
        source_hash = hashlib.sha256(data.SOURCE.read_bytes()).hexdigest()
        with tempfile.TemporaryDirectory(prefix=".test-", dir=data.GENERATED) as directory:
            with patch.object(data, "GENERATED", Path(directory)), redirect_stdout(io.StringIO()), redirect_stderr(io.StringIO()):
                self.assertEqual(0, data.main([]))
                first = {p.name: p.read_bytes() for p in Path(directory).glob("*.json")}
                self.assertEqual(0, data.main([]))
                self.assertEqual(first, {p.name: p.read_bytes() for p in Path(directory).glob("*.json")})
                self.assertEqual(0, data.main(["--check"]))
                self.tables["Stages"][0]["OpeningMoves"] = "7"
                with patch.object(data, "read_source", return_value=self.tables):
                    self.assertEqual(1, data.main([]))
                self.assertEqual(first, {p.name: p.read_bytes() for p in Path(directory).glob("*.json")})
                (Path(directory) / "ads.json").write_text("{}", encoding="utf-8")
                self.assertEqual(1, data.main(["--check"]))
        self.assertEqual(source_hash, hashlib.sha256(data.SOURCE.read_bytes()).hexdigest())

    def test_reject_unavoidable_next_reply_loss(self):
        # AI has bottom-row stones 1,2,3; both ends 0 and 4 win, and player has no immediate win.
        moves = "6,1,6,2,5,3"
        board, side = data.replay_opening(moves, 1)
        self.assertFalse(data.can_survive_next_reply(board, side))
        self.assert_invalid("Stages", 0, "OpeningMoves", moves, "every legal move loses")
        self.assertFalse(audit.four(0))

    def test_bitboard_matches_independent_whole_board_win_windows(self):
        for stage in self.tables["Stages"]:
            sequence = [int(x) for x in (stage["OpeningMoves"] or "").split(",") if x]
            stacks, side = independent_board(sequence, stage["Starter"])
            own, enemy, heights = audit.position(stage)
            for player, bits in ((side, own), (-side, enemy)):
                expected = immediate_wins(stacks, player)
                actual = [col for col in range(7) if heights[col] < 6
                          and audit.four(bits | (1 << (col * 7 + heights[col])))]
                self.assertEqual(expected, actual)

    def test_bounded_search_matches_unpruned_short_tree(self):
        def oracle(stacks, side, depth):
            if depth == 0:
                return 0
            if immediate_wins(stacks, side):
                return 1
            outcomes = []
            for col in range(7):
                if len(stacks[col]) == 6:
                    continue
                child = copy.deepcopy(stacks)
                child[col].append(side)
                outcomes.append(-oracle(child, -side, depth - 1))
            return max(outcomes, default=0)
        for index in (0, 2, 8, 15, 25):
            stage = self.tables["Stages"][index]
            sequence = [int(x) for x in (stage["OpeningMoves"] or "").split(",") if x]
            stacks, side = independent_board(sequence, stage["Starter"])
            self.assertEqual(oracle(stacks, side, 3), audit.solve(*audit.position(stage), 3))
        audit.solve.cache_clear()

    def test_one_move_wins_for_either_side_are_rejected(self):
        # Both positions are legal and unfinished, but already contain a winning next move.
        for moves in ('4,2,5,1,6,2', '6,0,6,1,5,2'):
            with self.subTest(moves=moves):
                self.tables = copy.deepcopy(self.original)
                self.assert_invalid('Stages', 0, 'OpeningMoves', moves, 'one-move win')

    def test_replace_failure_rolls_back_complete_bundle(self):
        old = data.validate(self.tables)
        changed = copy.deepcopy(old)
        changed["stages.json"]["stages"][0]["titleKey"] = "stage_replaced"
        changed["strings.json"]["strings"].append({"key": "stage_replaced", "ko": "교체 시험"})
        actual_replace = data.os.replace
        for fail_at in range(1, len(old) + 1):
            with self.subTest(fail_at=fail_at), tempfile.TemporaryDirectory(prefix=".test-", dir=data.GENERATED) as directory:
                with patch.object(data, "GENERATED", Path(directory)):
                    data.write_outputs(old)
                    before = {name: (Path(directory) / name).read_bytes() for name in old}
                    calls = 0
                    def fail_one(source, destination):
                        nonlocal calls
                        calls += 1
                        if calls == fail_at:
                            raise OSError("injected publication failure")
                        return actual_replace(source, destination)
                    with patch.object(data.os, "replace", side_effect=fail_one):
                        with self.assertRaisesRegex(OSError, "injected"):
                            data.write_outputs(changed)
                    self.assertEqual(before, {name: (Path(directory) / name).read_bytes() for name in old})
                    self.assertFalse((Path(directory) / data.INCOMPLETE_MARKER).exists())
                    self.assertEqual(len(old), len(list(Path(directory).iterdir())))

    def test_failed_first_publication_restores_file_absence(self):
        payloads = data.validate(self.tables)
        actual_replace = data.os.replace
        with tempfile.TemporaryDirectory(prefix=".test-", dir=data.GENERATED) as directory:
            with patch.object(data, "GENERATED", Path(directory)):
                calls = 0
                def fail_second(source, destination):
                    nonlocal calls
                    calls += 1
                    if calls == 2:
                        raise OSError("injected first-bundle failure")
                    return actual_replace(source, destination)
                with patch.object(data.os, "replace", side_effect=fail_second), self.assertRaises(OSError):
                    data.write_outputs(payloads)
                self.assertEqual([], list(Path(directory).iterdir()))

    def test_rollback_failure_retains_marker_backups_and_blocks_check(self):
        old = data.validate(self.tables)
        changed = copy.deepcopy(old)
        changed["stages.json"]["stages"][0]["titleKey"] = "stage_replaced"
        actual_replace = data.os.replace
        with tempfile.TemporaryDirectory(prefix=".test-", dir=data.GENERATED) as directory:
            with patch.object(data, "GENERATED", Path(directory)):
                data.write_outputs(old)
                calls = 0
                def fail_publication_and_rollback(source, destination):
                    nonlocal calls
                    calls += 1
                    if calls in (2, 3):
                        raise OSError("injected write/rollback failure")
                    return actual_replace(source, destination)
                with patch.object(data.os, "replace", side_effect=fail_publication_and_rollback):
                    with self.assertRaisesRegex(OSError, "rollback is incomplete"):
                        data.write_outputs(changed)
                marker = Path(directory) / data.INCOMPLETE_MARKER
                self.assertTrue(marker.exists())
                record = json.loads(marker.read_text(encoding="utf-8"))
                backups = Path(directory) / record["backupDirectory"]
                self.assertEqual(data.serialize(old)["stages.json"], (backups / "previous-stages.json").read_bytes())
                with redirect_stderr(io.StringIO()):
                    self.assertEqual(1, data.main(["--check"]))
                with self.assertRaisesRegex(OSError, "Incomplete export"):
                    data.write_outputs(old)


if __name__ == "__main__":
    unittest.main()
