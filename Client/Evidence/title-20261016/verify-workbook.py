from pathlib import Path
from copy import copy
import json
import openpyxl
import shutil

root=Path(__file__).resolve().parents[3]
e=Path(__file__).parent
before=openpyxl.load_workbook(e/'GameData-before.xlsx')
after=openpyxl.load_workbook(e/'GameData-candidate.xlsx')
assert before.sheetnames==after.sheetnames
changed=[]
for a,b in zip(before,after):
    assert (a.max_row,a.max_column)==(b.max_row,b.max_column)
    assert str(a.freeze_panes)==str(b.freeze_panes)
    assert str(a.merged_cells)==str(b.merged_cells)
    assert a.data_validations==b.data_validations
    assert a.auto_filter==b.auto_filter
    for row in a:
        for x in row:
            y=b[x.coordinate]
            if x.value!=y.value:changed.append({'sheet':a.title,'cell':x.coordinate,'before':x.value,'after':y.value})
            for attr in ('font','fill','border','alignment','protection','number_format'):
                assert copy(getattr(x,attr))==copy(getattr(y,attr)),(a.title,x.coordinate,attr)
    for key,value in a.column_dimensions.items():
        assert value.width==b.column_dimensions[key].width
    for key,value in a.row_dimensions.items():
        assert value.height==b.row_dimensions[key].height
assert len(changed)==9 and {c['sheet'] for c in changed}=={'Strings'}
assert {c['cell'] for c in changed}=={'C35','D35','E35','C63','E63','C64','E64','C127','E127'}
source=root/'GameData/GameData_MVP.xlsx'
assert source.read_bytes()==(e/'GameData-before.xlsx').read_bytes(),'Workbook changed concurrently.'
(e/'workbook-verification.json').write_text(json.dumps({'changedCells':changed,'otherCellsAndStylesPreserved':True,'sheetsPreserved':True},ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
shutil.copyfile(e/'GameData-candidate.xlsx',source)
print('Applied 8 translated title cells and one title-status note; other values and styles preserved.')
