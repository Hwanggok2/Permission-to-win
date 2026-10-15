"""Read-only comparison: preserve every cell/feature except the authored E/F/G profile values."""
import json
from pathlib import Path
import openpyxl
folder=Path(__file__).resolve().parent
before=openpyxl.load_workbook(folder/'GameData-before.xlsx')
after=openpyxl.load_workbook(folder/'GameData-candidate.xlsx')
profiles=json.loads((folder/'final-profiles.json').read_text(encoding='utf-8-sig'))
assert before.sheetnames==after.sheetnames
changes=[]
for a,b in zip(before,after):
    assert a.max_row==b.max_row and a.max_column==b.max_column
    assert str(a.merged_cells)==str(b.merged_cells)
    assert str(a.data_validations)==str(b.data_validations)
    assert a.freeze_panes==b.freeze_panes
    assert a.auto_filter==b.auto_filter
    assert list(a.tables)==list(b.tables)
    assert len(a.conditional_formatting)==len(b.conditional_formatting)
    for key,dim in a.row_dimensions.items():
        assert (dim.height,dim.hidden)==(b.row_dimensions[key].height,b.row_dimensions[key].hidden)
    for key,dim in a.column_dimensions.items():
        assert (dim.width,dim.hidden)==(b.column_dimensions[key].width,b.column_dimensions[key].hidden)
    for row in a:
        for cell in row:
            other=b[cell.coordinate]
            assert cell._style==other._style,(a.title,cell.coordinate,'style')
            expected=cell.value
            if a.title=='Stages' and 5<=cell.row<=34 and 5<=cell.column<=7:
                expected=profiles[(cell.row-5)//2][{5:'aiDepth',6:'aiNodeBudget',7:'aiRandomPermille'}[cell.column]]
            assert other.value==expected,(a.title,cell.coordinate,expected,other.value)
            if cell.value!=other.value:changes.append(dict(sheet=a.title,cell=cell.coordinate,before=cell.value,after=other.value))
assert changes
(folder/'workbook-changes.json').write_text(json.dumps(dict(changedCells=changes,otherValuesAndStylesPreserved=True,sheetFeaturesPreserved=True),ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
print(f'Workbook verified: {len(changes)} profile cells changed; other values, styles and sheet features preserved.')
