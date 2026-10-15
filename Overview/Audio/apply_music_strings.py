"""Preserve workbook formatting while publishing the two Artifact Tool-authored rows."""
from pathlib import Path
from zipfile import ZipFile
import copy
import hashlib
import json
import re
import xml.etree.ElementTree as ET
import openpyxl  # Read-only comparison; authoring is prepare_music_strings.mjs.

ROOT = Path(__file__).resolve().parents[2]
EVIDENCE = ROOT / 'Overview/Audio/evidence/integration-20261015'
SOURCE = ROOT / 'GameData/GameData_MVP.xlsx'
CANDIDATE = EVIDENCE / 'workbook-candidate.xlsx'
NS = 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'
ET.register_namespace('x', NS)

before = openpyxl.load_workbook(SOURCE)
candidate = openpyxl.load_workbook(CANDIDATE)
assert before.sheetnames == candidate.sheetnames
for sheet in before:
    for row in sheet:
        for cell in row:
            assert cell.value == candidate[sheet.title][cell.coordinate].value
assert before['Strings'].max_row == 422
assert candidate['Strings'].max_row == 424
assert candidate['Strings']['B423'].value == 'settings_music_on'
assert candidate['Strings']['B424'].value == 'settings_music_off'

with ZipFile(SOURCE) as archive:
    originals = {info.filename: (info, archive.read(info.filename)) for info in archive.infolist()}
sheet_path = 'xl/worksheets/sheet3.xml'
xml = originals[sheet_path][1].decode('utf-8')
template = ET.fromstring(xml).find(f'{{{NS}}}sheetData/{{{NS}}}row[@r="422"]')
rows = []
for number in (423, 424):
    row = copy.deepcopy(template)
    row.set('r', str(number))
    for cell in row:
        column = re.match(r'[A-Z]+', cell.get('r')).group(0)
        cell.set('r', column + str(number))
        cell.set('t', 'inlineStr')
        for child in list(cell):
            cell.remove(child)
        ET.SubElement(ET.SubElement(cell, f'{{{NS}}}is'), f'{{{NS}}}t').text = str(candidate['Strings'][column + str(number)].value)
    rows.append(ET.tostring(row, encoding='unicode'))
closing = '</x:sheetData>' if '</x:sheetData>' in xml else '</sheetData>'
xml = xml.replace(closing, ''.join(rows) + closing)
xml = re.sub(r'(<dimension\b[^>]*\bref=")A1:E422(")', r'\g<1>A1:E424\2', xml)
final = EVIDENCE / 'workbook-preserved.xlsx'
with ZipFile(final, 'w') as archive:
    for name, (info, content) in originals.items():
        archive.writestr(info, xml.encode('utf-8') if name == sheet_path else content)

after = openpyxl.load_workbook(final)
assert after['Strings'].max_row == 424
for sheet in before:
    for row in sheet:
        for cell in row:
            target = after[sheet.title][cell.coordinate]
            assert cell.value == target.value and cell._style == target._style
with ZipFile(final) as archive:
    changed_parts = [name for name, (_, content) in originals.items() if archive.read(name) != content]
assert changed_parts == [sheet_path]
SOURCE.write_bytes(final.read_bytes())
report = {
    'authoring': 'Artifact Tool, two new Strings rows; original OOXML parts preserved during publication because full export changed unrelated styles',
    'added_range': 'Strings!A423:E424', 'added_keys': ['settings_music_on', 'settings_music_off'],
    'unchanged_existing_cells': True, 'unchanged_existing_cell_styles': True,
    'changed_zip_parts': changed_parts, 'sha256': hashlib.sha256(SOURCE.read_bytes()).hexdigest(),
}
(EVIDENCE / 'workbook-validation.json').write_text(json.dumps(report, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
print(json.dumps(report, ensure_ascii=True))
